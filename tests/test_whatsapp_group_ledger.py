from pathlib import Path

import pytest


def test_extract_ledger_fields():
    from plugins import whatsapp_group_ledger as ledger

    text = """
    Payment: Rs 4500
    Bill No: AT-2048
    Sender Name: Riya Shah
    """

    record = ledger._extract_ledger_record(text)
    assert record["payment_amount"] == 4500.0
    assert record["bill_no"] == "AT-2048"
    assert record["sender_name"] == "Riya Shah"


def test_duplicate_hash_is_detected(tmp_path):
    from plugins import whatsapp_group_ledger as ledger

    workbook = tmp_path / "ledger.xlsx"
    ledger._write_ledger_excel(workbook, [
        {"date": "2026-10-10", "group_name": "Demo", "hash": "abc123"}
    ])

    assert ledger._existing_hashes(workbook) == {"abc123"}
    assert ledger._is_duplicate(workbook, "abc123") is True


def test_whatsapp_details_are_saved_and_same_bill_is_skipped(tmp_path):
    from openpyxl import load_workbook
    from plugins import whatsapp_group_ledger as ledger

    workbook = tmp_path / "payments.xlsx"
    details = {
        "group_name": "ONLINE PAY CONF",
        "output_xlsx": str(workbook),
        "sender_name": "Riya Shah",
        "bill_no": "AT-2048",
        "payment_amount": 4500,
        "date": "2026-10-10",
    }

    saved = ledger._process_group_images(details)
    assert "Saved WhatsApp payment record" in saved
    assert workbook.exists()

    details["date"] = "2026-10-11"
    duplicate = ledger._process_group_images(details)
    assert "Duplicate payment record skipped" in duplicate

    wb = load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Ledger"]
    assert ws.max_row == 2
    assert ws.cell(2, 2).value == "ONLINE PAY CONF"
    assert ws.cell(2, 3).value == "Riya Shah"
    assert ws.cell(2, 4).value == "AT-2048"
    assert ws.cell(2, 5).value == 4500
    wb.close()


def test_direct_save_creates_configured_sheet_headers(tmp_path):
    from openpyxl import Workbook, load_workbook
    from plugins import whatsapp_group_ledger as ledger

    workbook = tmp_path / "existing.xlsx"
    wb = Workbook()
    wb.active.title = "Notes"
    wb.save(workbook)
    wb.close()

    result = ledger._process_group_images({
        "group_name": "Payments",
        "output_xlsx": str(workbook),
        "sheet_name": "WhatsApp Ledger",
        "record_text": "Payment: Rs 1200\nBill No: B-7\nSender Name: Aman",
    })

    assert "Saved WhatsApp payment record" in result
    wb = load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["WhatsApp Ledger"]
    assert ws.cell(1, 1).value == "date"
    assert ws.cell(2, 5).value == 1200
    wb.close()


def test_direct_save_uses_documents_default_when_path_missing(tmp_path, monkeypatch):
    from plugins import whatsapp_group_ledger as ledger

    workbook = tmp_path / "Documents" / "WhatsApp_Payment_Ledger.xlsx"
    monkeypatch.setattr(ledger, "_default_output_path", lambda: workbook)
    result = ledger._process_group_images({
        "group_name": "Payments",
        "record_text": "Payment: Rs 1200\nBill No: B-8\nSender Name: Aman",
    })

    assert "Saved WhatsApp payment record" in result
    assert workbook.exists()
